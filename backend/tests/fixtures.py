"""Synthetic .xlsx fixture builders for parser tests. No real customer data."""

from __future__ import annotations

import io
from datetime import date, datetime
from openpyxl import Workbook

# ---- GreytHR ---------------------------------------------------------------


def build_greythr_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = [
        "Sl No",
        "Employee No",
        "Name",
        "Manager No",
        "Manager Name",
        "Leave Type",
        "Transaction Type",
        "Posted Date",
        "From Date",
        "To Date",
        "Days",
        "Expire Date",
        "Reason",
        "Remarks",
    ]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def date_to_excel_serial(d: date) -> int:
    """Inverse of utils.date_utils.excel_serial_to_date."""
    return (datetime(d.year, d.month, d.day) - datetime(1899, 12, 30)).days


# ---- Biometric (block format per PRD §14.2) --------------------------------


def build_biometric_xlsx(blocks: list[dict]) -> bytes:
    """Build the report-style block format.

    Each block:
        {
            "emp_code": "E001",
            "emp_name": "Alice Doe",
            "days": [
                {
                    "att_date": date(2026,4,1),
                    "in_time": "10:15",
                    "out_time": "19:00",
                    "shift": "UB",
                    "sched_in": "10:00",
                    "sched_out": "19:00",
                    "work_dur": "08:45",
                    "ot": "00:00",
                    "total_dur": "08:45",
                    "late_by": "00:15",
                    "early_by": "00:00",
                    "status": "Present",
                    "punch": "10:15-19:00",
                },
                ...
            ],
        }
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Three header rows at the very top (skipped by the parser).
    ws.append(["", "Daily Attendance Report (Detailed Summary Report)"])
    ws.append(["", "01-Apr-2026 to 30-Apr-2026"])
    ws.append(["", "Company:", "Shorthills GGn", "", "", "", "Printed On: 01-May-2026"])

    for blk in blocks:
        # EMPLOYEE MARKER row: col B = "Emp Code:", col C = code, col E = "Employee Name :", then name.
        row = ["", "Emp Code:", blk["emp_code"], "", "Employee Name :", blk["emp_name"]]
        ws.append(row)
        # Column header row (skipped via skip_next_row).
        ws.append(
            [
                "",
                "Att. Date",
                "InTime",
                "OutTime",
                "Shift",
                "",
                "Sched In",
                "Sched Out",
                "",
                "Work Duration",
                "OT",
                "Total Duration",
                "LateBy",
                "EarlyGoingBy",
                "Status",
                "Punch",
            ]
        )
        for d in blk["days"]:
            ws.append(
                [
                    "",
                    d["att_date"].strftime("%d-%b-%Y"),
                    d.get("in_time"),
                    d.get("out_time"),
                    d.get("shift"),
                    "",
                    d.get("sched_in"),
                    d.get("sched_out"),
                    "",
                    d.get("work_dur"),
                    d.get("ot"),
                    d.get("total_dur"),
                    d.get("late_by"),
                    d.get("early_by"),
                    d.get("status"),
                    d.get("punch"),
                ]
            )
        # SUMMARY row ends the block.
        ws.append(["", "Total Duration=176:30", "", "OT=00:00"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---- Roster ----------------------------------------------------------------

ROSTER_COLUMNS = [
    "Name",
    "Project",
    "Client",
    "In team Role",
    "Billing status",
    "Working Model",
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
    "Sunday",
    "Intime window onpen till",  # typo verbatim
    "Shift Time (Day)",
    "Shift Time (Night)",
    "Comments",
    "WFH Credits",
]


def build_roster_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(ROSTER_COLUMNS)
    for r in rows:
        ws.append([r.get(c) for c in ROSTER_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
